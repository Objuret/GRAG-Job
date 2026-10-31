"""The finished data, for giving away: one workbook, values only. Sheets: results (the three
arms' questions in one table), run data (what is the same for the arms once, then a column per
arm), arm construction (a column per arm), key.

    python tools/export_finished_data.py [--out finished_data]

Needs openpyxl. Reads the final run folders, one per arm and question set (RUNS), the questions
file they all read (data/questions.jsonl, checked against each run's hash of it), the files of the
question sets (SETS, which together are HERB's answerable questions, each once) and the graph's
build cost (test/graph/build_cost.json), and writes full-grag-data.xlsx. It prints which question sets
each arm's runs cover and names a set an arm has no run on. A run split over the laptop and the
desktop is read with its halves for its environment. Every run's time is its questions' times added,
its judging's time its judge calls' and embedding requests' times added; an arm's is its runs' added.
A model call counts by its answering try's server time; the run folders' corrections give the
answer and first calls past the login's limit, and the search times, their values (model-call-time,
search-time in each folder's corrections.json). A run that read its interpreter's answers
and its query embeddings from an earlier run takes those values from that run (MADE_IN); the artefact's
run on the 715 takes its query embeddings from a run of the same questions without answers that embedded
every question's texts in one pass, on the graphics card (EMBEDDED_IN). Two scores
judged after a run, on its answers with its judge, are read from the folder beside the run
(REJUDGED), and their judging is added to the run's. Since HERB's employee and customer records are
in every arm's corpus, a run without answers beside each final run (RETRIEVED) gives every question's
search over the index that holds them, and tells, text by text, which questions are now handed other
texts; those were answered and judged again, an arm's in one run (REDONE), and their rows, scores and
judging come from it, every other question's from the final run. A column of a run is a measure he named
(times, tokens, characters, runtime, the RAGAS scores) or an id of the dataset, the evaluation or
the run; the construction holds what built the index that each final run read, one model at a
time, and nothing about how it was built. Nothing that belongs to this machine or this project
goes in: no path, host, user, session, repository file, commit, note or reason. The data sheets
come first; the last sheet is the key.
"""
import argparse
import asyncio
import hashlib
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import correct_missing_judge_legs as judge_records
from compare_delivered import delivered

ROOT = Path(__file__).resolve().parent.parent
# an arm's final runs, one per question set; their rows go into the same sheets (his 2026-10-25
# "the fucking point is to have a full ragas-herb test done finally")
RUNS = {"lucene": ["lucene__gold100__cb72000__20261008T234209Z",
                   "lucene__answerable715__cb72000__20261023T112745Z"],
        "vector": ["vector__gold100__cb72000__20261017T061242Z",
                   "vector__answerable715__cb72000__20261024T004317Z"],
        "artefact": ["artefact_v4__gold100__cb72000__20261023T021439Z",
                     "artefact_v4__answerable715__cb72000__20261025T211532Z"]}
# HERB's answerable questions, each in one of these files of ids; a run's set is the file whose
# sha256 its manifest names as its ids file
SETS = {"gold100": "data/gold100.jsonl", "answerable715": "data/answerable715.jsonl"}
# the artefact's run on the hundred read its interpreter's answers and its query embeddings, the same
# on every question, from this run (his 2026-10-22 "it's the same as the fucking "old run's interpreter"")
MADE_IN = {"artefact_v4__gold100__cb72000__20261023T021439Z": "artefact_v4__gold100__cb72000__20261020T041923Z"}
# the artefact's run on the 715 takes its query embeddings, and only those, from a run that embedded every
# question's texts in one pass, as MADE_IN's run did for the hundred: the same questions run again without
# answers, each first call read from the run's own store, the embedding cache empty, on the graphics card.
# Starts of the run that the desktop's crashes ended had embedded some questions' texts, and their records
# went with them
EMBEDDED_IN = {"artefact_v4__answerable715__cb72000__20261025T211532Z":
               "artefact_v4__answerable715__cb72000__20261027_query_embeddings"}
# in EMBEDDED_IN's run: each text's vector against the run's own vector of the same text, by cosine
VECTOR_CHECK = "query_vectors_against_the_run.json"
# scores judged after a run, on its answers with its judge, by prod/run.py --rejudge into the folder
# beside the run, named by the run, this and the set (his 2026-10-23 "just toggle them and run them
# on all 100 for all 3 arms"); a run that judged every score itself has no such folder
REJUDGED = "__j-claude-haiku-4-5-20251001__"
REJUDGED_SCORES = ("factual_correctness", "context_entity_recall")
# beside each final run, the run of its questions without answers over the corpus with HERB's employee and
# customer records (his 2026-10-30 "ok, so, we add them back in then"): which of its questions are now handed
# other texts (compared text by text, tools/compare_delivered.py), the index the arm reads now, and each
# question's search over it, measured alone after the run, the six runs one after the other on a quiet desktop, one
# method and one condition for every question and arm
RETRIEVED = {"lucene__gold100__cb72000__20261008T234209Z": "lucene__gold100__cb72000__20261031_metadata_search_alone",
             "lucene__answerable715__cb72000__20261023T112745Z":
                 "lucene__answerable715__cb72000__20261031_metadata_search_alone",
             "vector__gold100__cb72000__20261017T061242Z": "vector__gold100__cb72000__20261031_metadata_search_alone",
             "vector__answerable715__cb72000__20261024T004317Z":
                 "vector__answerable715__cb72000__20261031_metadata_search_alone",
             "artefact_v4__gold100__cb72000__20261023T021439Z":
                 "artefact_v4__gold100__cb72000__20261031_metadata_search_alone",
             "artefact_v4__answerable715__cb72000__20261025T211532Z":
                 "artefact_v4__answerable715__cb72000__20261031_metadata_search_alone"}
# the questions now handed other texts, answered and judged again, an arm's in one run over both sets (his
# 2026-10-31 "fair enough" to that play, "go ahead and do this "for real""); every other question keeps its
# final run's answer, scores and judging
REDONE = {"lucene": "lucene__metadata_redo__cb72000__20261031",
          "vector": "vector__metadata_redo__cb72000__20261031",
          "artefact": "artefact_v4__metadata_redo__cb72000__20261031"}
DATASET, LICENCE = "HERB (Salesforce AI Research)", "CC BY-NC 4.0"
# the sheets' names
PER_QUESTION, PER_RUN, CONSTRUCTION, KEY = "results", "run data", "arm construction", "key"

# (field, meaning, unit): the data dictionary is written from these lists. A score's meaning is
# RAGAS's own definition (docs.ragas.io, available metrics) with the setting RAGAS 0.4.3 used.
SCORES = [
    ("faithfulness", "RAGAS Faithfulness. Claims in the answer supported by the retrieved context / all claims "
                     "in the answer.", "0-1"),
    ("answer_correctness", "RAGAS AnswerCorrectness. Weighted average of factual similarity to the reference "
                           "answer (weight 0.75) and semantic_similarity (weight 0.25). Reference answer: HERB's answer.", "0-1"),
    ("factual_correctness", "RAGAS FactualCorrectness. F1 of precision, TP / (TP + FP), and recall, TP / (TP + FN). "
                            "TP: claims in the answer present in the reference answer. FP: claims in the answer not "
                            "present in the reference answer. FN: claims in the reference answer not present in the "
                            "answer. Claims and their presence: the judge's, atomicity and coverage low. Rounded to "
                            "two decimals.", "0-1"),
    ("context_recall_id", "RAGAS IDBasedContextRecall. Reference context ids found in the retrieved context ids / "
                          "all reference context ids. Reference context ids: HERB's evidence record ids. Retrieved context ids: "
                          "the ids of the records whose whole text was given to the answer model.", "0-1"),
    ("context_precision_id", "RAGAS IDBasedContextPrecision. Retrieved context ids found in the reference context "
                             "ids / all retrieved context ids. Retrieved context ids: the ids of the records whose whole text "
                             "was given to the answer model.", "0-1"),
    ("context_recall_llm", "RAGAS LLMContextRecall. Claims in the reference answer supported by the retrieved "
                           "context / all claims in the reference answer.", "0-1"),
    ("context_entity_recall", "RAGAS ContextEntityRecall. Entities common to the reference answer and the retrieved "
                              "context / all entities in the reference answer. Entities: the judge's.", "0-1"),
    ("context_recall_nonllm", "RAGAS NonLLMContextRecall. Relevant contexts retrieved / all reference contexts. "
                              "Relevant: Levenshtein similarity to a reference context of more than 0.5. Reference context: the "
                              "content of a HERB evidence record: a message's text, a document's content and feedback, a "
                              "transcript, a chat's text, a link's description, a pull request's summary and review "
                              "comments. Retrieved contexts: the content of each record passed to the answer model.",
     "0-1"),
    ("context_precision_nonllm", "RAGAS NonLLMContextPrecisionWithReference. Context precision@K: mean of "
                                 "precision@k over the ranks k of the relevant retrieved contexts. Relevant: "
                                 "Levenshtein similarity to a reference context of 0.5 or more. Reference context: the content "
                                 "of a HERB evidence record: a message's text, a document's content and feedback, a transcript, "
                                 "a chat's text, a link's description, a pull request's summary and review comments. "
                                 "Retrieved contexts: the content of each record passed to the answer model, ranked in "
                                 "the arm's order.", "0-1"),
    ("semantic_similarity", "RAGAS SemanticSimilarity. Cosine similarity of the embeddings of the answer and the "
                            "reference answer.", "0-1"),
    ("string_similarity", "RAGAS NonLLMStringSimilarity. Levenshtein similarity of the answer and the reference "
                          "answer.", "0-1"),
    ("bleu", "RAGAS BleuScore. BLEU of the answer against the reference answer.", "0-1"),
    ("rouge", "RAGAS RougeScore. ROUGE-L F-measure of the answer against the reference answer.", "0-1"),
    ("chrf", "RAGAS ChrfScore. chrF of the answer against the reference answer.", "0-1"),
    ("exact_match", "RAGAS ExactMatch. 1 if the answer is the same as the reference answer, else 0.", "0 or 1"),
    ("string_presence", "RAGAS StringPresence. 1 if the answer contains the reference answer, else 0.", "0 or 1"),
]
QUESTIONS = [
    ("arm", "Retrieval method.", ""),
    ("question_id", "Question id: product, a (answerable), index of the question in the product's answerable "
                    "questions in HERB, from 0.", ""),
    ("question_type", "HERB question type.", ""),
    ("question", "Question, from HERB.", ""),
    ("reference_answer", "Reference answer: HERB's answer to the question, as RAGAS was given it (its parts "
                         "joined by spaces).", ""),
    ("reference_context_ids", "Reference context ids: the ids of HERB's evidence records for the question, "
                              "joined by ', '.", ""),
    ("answer", "Answer of the answer model.", ""),
    *SCORES,
    ("string_presence_per_id", "RAGAS StringPresence run on each item of HERB's reference list (one employee ID, PR "
                               "link or URL per item) instead of on the whole reference answer: the share of the items "
                               "the answer contains. Person, PR and URL questions; empty for content and company "
                               "questions.", "0-1"),
    ("string_presence_per_id_precision", "RAGAS StringPresence run the other way round: each distinct ID or link in "
                                         "the answer of the forms that the items of HERB's reference list have, looked "
                                         "for in the reference list (its items joined by spaces): the share of them the "
                                         "list contains. The forms are our own rule: an employee ID (eid_ and eight hex "
                                         "digits), a customer ID (CUST- and four digits), a PR link (a github.com pull "
                                         "link), a URL (any other link that starts with http:// or https://). Person, PR and "
                                         "URL questions; empty for content and company questions, and empty when the "
                                         "answer gives no ID or link of those forms.", "0-1"),
    ("string_presence_per_id_f1", "F1 of string_presence_per_id_precision (P) and string_presence_per_id (R): "
                                  "2PR / (P + R); 0 when the answer gives no ID or link of the forms of "
                                  "string_presence_per_id_precision, or when P + R is 0. The forms are our own rule. "
                                  "Person, PR and URL questions; empty for content and company questions.", "0-1"),
    ("search_seconds", "Retrieval time, without the interpreter and the query embedding, measured with one question "
                       "at a time on one machine for every arm, never as the first question of a process: for lucene "
                       "and vector the ranking, for the artefact the reading of its stored first-call answer and query "
                       "vectors, the ranking and its hand-over of the records up to character_limit.", "seconds"),
    ("interpreter_seconds", "Time of the interpreter call's answering try: the answer model, called before "
                            "retrieval, measured and limited as answer_seconds, with the relation fitted on the "
                            "interpreter calls. Artefact only.", "seconds"),
    ("query_embedding_seconds", "Time to embed the query's texts on the graphics card, one request after the other "
                                "in a warm process, without waits for the embedder. Vector and artefact only. The "
                                "artefact keeps each text's vector: a text that an earlier question of the same run "
                                "embedded is read with no request and counts with that question.", "seconds"),
    ("answer_seconds", "Time of the answer model call's answering try: the Claude CLI's duration_api_ms, its time in "
                       "requests to the model. For a call started while the backlog of answer-model input tokens of all "
                       "answer-model calls in the run folders, test runs included (added at each call's start, lowered "
                       "by 16,500 tokens per second), was 886,127 tokens or more, the fitted backlog at which the "
                       "seconds per output token rise: a + b x output tokens, "
                       "fitted on the same arm's calls in its final runs started below it. A call in which the CLI "
                       "did not accept the answer's form and the model wrote the answer again counts by its first request: "
                       "duration_api_ms times the first request's share of the requests' times in the CLI's transcript "
                       "of the call, or the relation at the first request's output tokens.",
     "seconds"),
    ("total_seconds", "Time of the question's steps added: search_seconds, interpreter_seconds, "
                      "query_embedding_seconds and answer_seconds. Without failed tries of model "
                      "calls.", "seconds"),
    ("interpreter_tokens_in", "Input tokens of the interpreter call: new input, cache creation and cache read "
                              "tokens added; of its first request, as answer_seconds says. Artefact only.", "tokens"),
    ("interpreter_tokens_out", "Output tokens of the interpreter call; of its first request, as answer_seconds says. "
                               "Artefact only.", "tokens"),
    ("query_embedding_tokens", "Tokens embedded for the query's texts, counted on each text with the 'query: ' "
                               "prefix and the tokenizer's special tokens. Vector and artefact only; for the "
                               "artefact, as query_embedding_seconds says which texts count.", "tokens"),
    ("retrieved_context_characters", "Characters of the retrieved context given to the answer model. At most "
                                     "character_limit.", "characters"),
    ("answer_tokens_in", "Input tokens of the answer model call: new input, cache creation and cache read tokens "
                         "added; of its first request, as answer_seconds says.", "tokens"),
    ("answer_tokens_out", "Output tokens of the answer model call; of its first request, as answer_seconds says.",
     "tokens"),
]
# the same for the three arms; the export stops if an arm differs. The dataset and its licence
# are their own words and have no entry in the data dictionary.
SETUP = [
    ("dataset", None, ""),
    ("dataset_licence", None, ""),
    ("character_limit", "Maximum characters of retrieved text given to the answer model per question.",
     "characters"),
    ("answer_model", "Language model of the answers and of the artefact's interpreter.", ""),
    ("judge_model", "Language model that RAGAS uses as judge.", ""),
    ("embedding_model", "Embedding model of the query embeddings and of the RAGAS scores that use embeddings.", ""),
]
RUN = [
    ("run_seconds", f"Time of the arm's runs over its questions in {PER_QUESTION}: their total_seconds added.",
     "seconds"),
    ("judge_tokens_in", "Input tokens of all RAGAS judge calls that answered: new input, cache creation and cache "
                        "read tokens added.", "tokens"),
    ("judge_tokens_out", "Output tokens of all RAGAS judge calls that answered.", "tokens"),
    ("judge_embedding_tokens", "Tokens embedded by RAGAS, counted on each text with its mode's prefix and the "
                               "tokenizer's special tokens.", "tokens"),
    ("judging_seconds", f"Time of the RAGAS evaluation over the arm's questions in {PER_QUESTION}: the Claude CLI's "
                        "duration_api_ms of every judge call that answered, as judge_tokens_in counts them (also in "
                        "cells that failed or that a later start judged again), and the encoding times of RAGAS's "
                        "embedding requests on the graphics card, added. Judge calls that no record holds (three of "
                        "the artefact, one of vector) are in the tokens and have no time.", "seconds"),
]
INDEX_KIND = {"lucene": ("BM25 full-text index", "corpus record"),
              "vector": ("dense vector index", "corpus record"),
              "artefact": ("graph database (Neo4j)", "chunk: a part of one corpus record, one record, or several records")}


def model_name(name):
    """A model's name without its publisher and without a snapshot date (his 2026-10-21 "fucking
    off with the nvidia tag", "20251001 for haiku? Stop beeing retarded")."""
    return None if name is None else re.sub(r"-\d{8}$", "", str(name).split("/")[-1])


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def rows(path):
    return [json.loads(x) for x in Path(path).read_text(encoding="utf-8").splitlines() if x.strip()]


def one(values, what):
    values = list(values)
    if any(v != values[0] for v in values):
        sys.exit(f"{what} differs between questions; this export reads one")
    return values[0]


def judging_seconds(calls, ids, what):
    """The judging's time: the time of all the things in it added, never a clock over the judging or
    over a question's judging (his 2026-10-22 "THE RUN TIME IS THE ADDED TIME FOR ALL THE THINGS IN THE
    FUCKING RUN", 2026-10-23 "wall time? this is not a thing"). A judge call that answered counts by its
    answering try's server time, the CLI's own duration_api_ms, without the program's start on our
    machine; one of RAGAS's embedding requests by its encoding time, encode_s, without its wait for the
    embedder's lock. How many calls ran at once changes neither. A try that did not answer counts with
    nothing (his 2026-10-21 "a model call's time is its answering try; the seconds of failed tries are
    out of the call, the question and the run"). Also the seconds of the tries so left out, by our clock."""
    total, failed, seen = 0.0, 0.0, set()
    for c in calls:
        if c.get("kind") != "chat":
            # RAGAS's own embedding requests, made for the run's answers and references before the
            # cells (phase prime): their record names no question
            total += float(c["encode_s"])
            continue
        if c.get("question_id") not in ids:
            continue
        seen.add(c["question_id"])
        tries = c.get("attempts") or []
        for n, t in enumerate(tries):
            if c.get("ok") and n == len(tries) - 1:
                server = (c.get("envelope") or {}).get("duration_api_ms")
                if not server:
                    sys.exit(f"{what}: a judge call that answered has no server time")
                total += server / 1000
            else:
                failed += float(t.get("seconds") or 0.0) + float(t.get("backoff_s") or 0.0)
    if seen != set(ids):
        sys.exit(f"{what}: not every question has its judge calls")
    return total, failed


def judge_cost(folder, ev, calls, cells, d):
    """judge_usage, checked against the scoring starts and the call records: every start in the call
    records has its entry in judge_legs, each entry's counts are those of its start's judge calls that
    answered (counted as tools/correct_missing_judge_legs.py counts them), judge_usage is the entries'
    sum, and the score cells' usage is their own calls'. And the tokens the evaluation embedded: the
    entries' embedding, checked against the embedder records."""
    counts = judge_records.COUNTS
    # a start stopped before any of its cells ended keeps the entry it wrote as it began (no settings, no usage) and
    # left no call record: no call of it answered, so it counts nothing (his 2026-10-21 "a model call's time is its
    # answering try"); a call record of a start with no entry still stops the export below
    legs = [leg for leg in ev["judge_legs"]
            if not (leg.get("begun_at") and leg.get("settings") is None and leg.get("usage") is None)]
    starts = [(leg.get("settings") or {}).get("started_at") for leg in legs]
    by_leg = {}
    for c in calls:
        by_leg.setdefault(c["eval_leg"], []).append(c)
    if not all(starts) or set(by_leg) - set(starts):
        sys.exit(f"{folder}: the scoring starts {sorted(set(by_leg) - set(starts))} are in eval_calls.jsonl "
                 f"and have no entry in judge_legs, or an entry names no start")
    found = (read(d / "judge_calls_without_record.json").get("starts") or {}
             if (d / "judge_calls_without_record.json").is_file() else {})
    for start, leg in zip(starts, legs):
        records = judge_records.usage_of(by_leg.get(start, []))
        xs = (found.get(start) or {}).get("named_by_no_record") or []
        if xs:
            # calls of a cell that had failed, which answered after the cell had ended: no record holds
            # them, the lane's own transcripts on the machine that made them name them
            records = dict(records, calls=records["calls"] + len(xs), attempts=records["attempts"] + len(xs),
                           tokens_in=records["tokens_in"] + sum(x["input_tokens"] + x["cache_creation_input_tokens"]
                                                                + x["cache_read_input_tokens"] for x in xs),
                           cached_input_tokens=records["cached_input_tokens"]
                           + sum(x["cache_read_input_tokens"] for x in xs),
                           tokens_out=records["tokens_out"] + sum(x["output_tokens"] for x in xs))
            print(f"{folder}: the entry of scoring start {start} holds {len(xs)} judge calls that no record holds, "
                  f"named in judge_calls_without_record.json")
        diff = judge_records.differs(leg["usage"], records, counts)
        # a written entry stays as recorded. A start that was cut short can have counted calls whose cell
        # wrote no record before it stopped, so its entry may hold more than its records; it may never hold
        # less, and no other entry may differ (the main session's decision of 2026-10-25, by his rule that
        # the data stays as recorded)
        if diff and (leg["settings"].get("interrupted")
                     and all(leg["usage"][k] >= records[k] for k in counts)):
            print(f"{folder}: the entry of scoring start {start}, which was cut short, holds more than its "
                  f"records (entry, records): {diff}; the entry stands")
        elif diff:
            sys.exit(f"{folder}: the entry of scoring start {start} and its judge call records differ "
                     f"(entry, records): {diff}")
    if judge_records.differs(judge_records.summed(legs), ev["judge_usage"], counts):
        sys.exit(f"{folder}: judge_usage is not the sum of judge_legs")
    own = {(qid, m, c["components"]["usage"]["eval_leg"]) for qid, q in cells.items() for m, c in q.items()}
    use = [c["components"]["usage"] for q in cells.values() for c in q.values()]
    mine = [c for c in calls if c.get("kind") == "chat" and (c.get("question_id"), c.get("score"), c["eval_leg"]) in own]
    if ((sum(u["chat_calls"] for u in use), sum(u["tokens_in"] for u in use), sum(u["tokens_out"] for u in use))
            != (len(mine), sum(int((c.get("usage") or {}).get("prompt_tokens") or 0) for c in mine),
                sum(int((c.get("usage") or {}).get("completion_tokens") or 0) for c in mine))
            or not all(c.get("ok") for c in mine)):
        sys.exit(f"{folder}: the scores' judge cost is not that of their own judge calls")
    embedded = sum(leg["settings"]["embedding"]["tokens_in"] for leg in legs)
    if embedded != judge_records.embedding_of(calls)["tokens_in"]:
        sys.exit(f"{folder}: the scoring starts' embedding tokens are not those of the embedder records")
    # the cost is every judge call that answered, the cells that failed or that a later start dropped
    # included: judge_usage (the main session's decision of 2026-10-25, by the key's "all RAGAS judge calls")
    return ev["judge_usage"], embedded


_REDONE = {}


def redone(arm):
    """The arm's run that answered and judged again the questions now handed other texts (REDONE), read once
    and checked as a final run is: every question answered and judged, the judge legs and their usage those of
    the call records. Its rows and scores by question, its call records, the questions a set took from it
    (used) and the lines of its embedding requests a set has counted (claimed), so that none counts twice."""
    if arm not in _REDONE:
        folder = REDONE[arm]
        d = ROOT / "output/k=chars" / folder
        man, ev = read(d / "run_manifest.json"), read(d / "eval_manifest.json")
        outputs = {o["id"]: o for o in rows(d / "arm_outputs.jsonl")}
        if (man["n_questions"], man["n_ran"], man["n_failed"]) != (len(outputs), len(outputs), 0):
            sys.exit(f"{folder}: not every question was answered")
        cells = {}
        for r in rows(d / "eval_results.jsonl"):
            cells.setdefault(r["question_id"], {})[r["metric"]] = r
        if set(cells) != set(outputs):
            sys.exit(f"{folder}: not every question it answered was judged")
        calls = rows(d / "eval_calls.jsonl")
        judge, _ = judge_cost(folder, ev, calls, cells, d)
        judging_seconds(calls, set(cells), folder)
        recorded = judge_records.usage_of(calls)
        if (judge["tokens_in"], judge["tokens_out"]) != (recorded["tokens_in"], recorded["tokens_out"]):
            sys.exit(f"{folder}: judge_usage holds judge calls that no record ties to a question")
        _REDONE[arm] = {"folder": folder, "man": man, "ev": ev, "rows": outputs, "cells": cells, "calls": calls,
                        "used": set(), "claimed": set()}
    return _REDONE[arm]


def judging_of(calls, ids, texts, what, claimed=frozenset()):
    """The judging of the questions ids in one folder's call records: their judge calls, each that answered by
    its answering try's server time and with its tokens, and every one of RAGAS's embedding requests whose text
    is one that their scores embed (texts), also when a later start made it again; a request already counted
    for another set's questions (claimed: its line) is not counted again. Returns the seconds, the seconds of
    the tries that did not answer and of their pauses, the judge calls' tokens in and out, the tokens embedded,
    and the lines of the requests counted."""
    chat = [c for c in calls if c.get("kind") == "chat" and c.get("question_id") in ids]
    seconds, failed = judging_seconds(chat, set(ids), what)
    use = judge_records.usage_of(chat)
    counted, embedded = set(), 0
    for n, c in enumerate(calls):
        if c.get("kind") == "chat":
            continue
        if c.get("kind") != "embed" or len(c.get("texts") or []) != 1:
            sys.exit(f"{what}: a record of the scoring is neither a judge call nor an embedding request of one text")
        if c["texts"][0] in texts and n not in claimed:
            seconds += float(c["encode_s"])
            embedded += int(c["tokens_in"])
            counted.add(n)
    return seconds, failed, use["tokens_in"], use["tokens_out"], embedded, counted


ID_TYPES = ("person", "pr", "url")  # HERB's question types whose reference list holds one ID or link per item
# the forms of an ID or a link, our own rule read from the corpus (the check of 2026-10-28); HERB's own
# evaluation has a model take the items out of the answer (code/evaluate.py)
URL = re.compile(r"https?://[^\s<>\"'()\[\]{}]+")
URL_TRAIL = ".,;:!?*`"  # with the bold and code marks the answers' markdown sets after a link
PR_LINK = re.compile(r"https://github\.com/[^/\s]+/[^/\s]+/pull/\d+")
EMPLOYEE_ID = re.compile(r"(?<![A-Za-z0-9_])eid_[0-9a-f]{8}(?![A-Za-z0-9])")
CUSTOMER_ID = re.compile(r"(?<![A-Za-z0-9])CUST-\d{4}(?![0-9])")


def ids_and_links(text):
    """The distinct IDs and links of a text, as (form, ID or link), in first-seen order: the links, less their
    trailing punctuation, a link that is a github.com pull link a PR link and any other a URL; then the employee
    and customer IDs in the text with its links masked."""
    found = []
    masked = text
    for m in URL.finditer(text):
        u = m.group(0).rstrip(URL_TRAIL)
        kind = "pr_link" if PR_LINK.fullmatch(u) else "other_url"
        found.append((kind, u))
        masked = masked[:m.start()] + " " * (m.end() - m.start()) + masked[m.end():]
    for m in EMPLOYEE_ID.finditer(masked):
        found.append(("employee_id", m.group(0)))
    for m in CUSTOMER_ID.finditer(masked):
        found.append(("customer_id", m.group(0)))
    seen, out = set(), []
    for k, i in found:
        if i not in seen:
            seen.add(i)
            out.append((k, i))
    return out


def string_presence(response, references):
    """RAGAS StringPresence of each reference in the response, 1 or 0."""
    from ragas.dataset_schema import SingleTurnSample
    from ragas.metrics import StringPresence
    metric = StringPresence()

    async def each():
        return [await metric.single_turn_ascore(SingleTurnSample(response=response, reference=str(r)))
                for r in references]
    return asyncio.run(each())


def string_presence_per_id(answer, items):
    """RAGAS StringPresence run on each item of HERB's reference list: the share of the items the answer holds."""
    hits = string_presence(answer, items)
    return sum(hits) / len(hits)


def string_presence_per_id_precision(answer, items):
    """RAGAS StringPresence run the other way round: each distinct ID or link of the answer, of the forms the
    reference items have, looked for in the items joined by spaces; the share of them found. None when the
    answer gives none of those forms."""
    forms = {k for item in items for k, _ in ids_and_links(str(item))}
    given = [i for k, i in ids_and_links(answer) if k in forms]
    if not given:
        return None
    hits = string_presence(" ".join(str(item) for item in items), given)
    return sum(hits) / len(hits)


def failed_tries_seconds(records):
    """The seconds of the model calls' tries that did not answer, and of the pauses after tries, as
    the call records hold them; an embedder request makes no try. A try marked machine_asleep was
    taken out of retry_s whole by the sleep correction (slept-try/1, the folder's corrections.json),
    so it counts here with no seconds; its pause stays."""
    return sum(sum(float(t.get("seconds") or 0.0) for t in c.get("attempts") or []
                   if t.get("outcome") != "ok" and not t.get("machine_asleep"))
               + sum(float(t.get("backoff_s") or 0.0) for t in c.get("attempts") or [])
               for c in records if c.get("kind") == "chat")


def interpreter_and_embeddings(r, vectors=True):
    """What the interpreter answered and what was embedded for the query, without how each was got
    (asked anew, or read from a store or a cache); without the vectors' sha256 when vectors is False
    (vectors made again on another device or in other batches differ in their last bits)."""
    i = r["meta"]["interpreter"]
    embedded = i["embeddings"] if vectors else {
        k: {f: v for f, v in g.items() if f != "embedding"} for k, g in i["embeddings"].items()}
    return ([{k: v for k, v in s.items() if k not in ("cache_hit", "offline", "asks")} for s in i["stages"]],
            embedded)


def halves(d, name):
    """A file of each half of a run split over two machines, as that half wrote it."""
    return [read(p) for p in sorted((d / "halves").glob(f"*/{name}"))]


def environment(d, man):
    """The run's environment; a split run keeps it in its halves, which must name one embedding model."""
    if man["environment"] is not None:
        return man["environment"]
    envs = [h["environment"] for h in halves(d, "run_manifest.json")]
    if not envs or len({e["embedder"]["model"] for e in envs}) != 1:
        sys.exit(f"{d.name}: the halves do not name one embedding model")
    return envs[0]


def rejudged(arm, d, cells, judge_model, set_name, keep):
    """REJUDGED_SCORES from the folder beside the run, which judged the run's own answers and holds the
    run's other scores as the run has them. Their cost is that of every judge call of the folder that
    answered, as the run's own judging is counted (judge_usage; the key's "all RAGAS judge calls"); their
    time their judge calls' server times added, as the run's. Only the questions that keep the run's answer
    (keep) count; the others were judged again (REDONE). A run with no such folder must have judged every
    score itself."""
    j = d.parent / (d.name + REJUDGED + set_name)
    if not j.exists():
        return {}, 0.0, 0, 0, 0.0
    if any(s in q for q in cells.values() for s in REJUDGED_SCORES):
        sys.exit(f"{arm}: {d.name} judged {REJUDGED_SCORES} itself and {j.name} judged them again")
    judged = lambda f: {r["id"]: (r["question"], r["answer"], r["contexts"], r["context_ids"])  # noqa: E731
                        for r in rows(f / "arm_outputs.jsonl")}
    if judged(j) != judged(d):
        sys.exit(f"{arm}: {j.name} judged other answers than the run's")
    added, kept = {}, set()
    for r in rows(j / "eval_results.jsonl"):
        if r["metric"] in REJUDGED_SCORES:
            added.setdefault(r["question_id"], {})[r["metric"]] = r
            kept.add((r["question_id"], r["metric"], r["components"]["usage"]["eval_leg"]))
        elif r != cells[r["question_id"]][r["metric"]]:
            sys.exit(f"{arm} {r['question_id']}: {r['metric']} in {j.name} is not the run's")
    every = rows(j / "eval_calls.jsonl")
    if any(c.get("score") not in REJUDGED_SCORES or c.get("question_id") not in cells for c in every):
        sys.exit(f"{arm}: {j.name} holds a call for another score or question than the re-judged ones")
    calls = [c for c in every if (c.get("question_id"), c.get("score"), c.get("eval_leg")) in kept]
    use = [c["components"]["usage"] for q in added.values() for c in q.values()]
    tokens = [c.get("usage") or {} for c in calls]
    if ((sum(u["chat_calls"] for u in use), sum(u["tokens_in"] for u in use), sum(u["tokens_out"] for u in use))
            != (len(calls), sum(int(t.get("prompt_tokens") or 0) for t in tokens),
                sum(int(t.get("completion_tokens") or 0) for t in tokens))
            or any(c["kind"] != "chat" or c["model"] != judge_model for c in every)):
        sys.exit(f"{arm}: the re-judged scores' cost is not their calls', or a call is not the run's judge's")
    # one basis for every judging: every judge call that answered (the main session's decision of
    # 2026-10-25, by the key's "all RAGAS judge calls"), of the questions that keep the run's answer
    judging_seconds(every, set(cells), f"{arm}: {j.name}")
    mine = [c for c in every if c.get("question_id") in keep]
    cost = judge_records.usage_of(mine)
    seconds, failed = judging_seconds(mine, set(keep), f"{arm}: {j.name}")
    return ({q: s for q, s in added.items() if q in keep}, seconds, cost["tokens_in"], cost["tokens_out"],
            failed)


def references():
    """Each question's row of the questions file, which holds HERB's reference answer and evidence record
    ids; every final run must have read this file as it is."""
    path = ROOT / "data/questions.jsonl"
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    for folder in [f for folders in RUNS.values() for f in folders] + list(RETRIEVED.values()) + list(REDONE.values()):
        if read(ROOT / "output/k=chars" / folder / "run_manifest.json")["inputs"]["questions_sha256"] != digest:
            sys.exit(f"{folder}: the run read another questions file than {path}")
    return {q["id"]: q for q in rows(path)}


def question_sets(refs):
    """Each set's question ids and, by the sha256 of its file, its name. Together the sets are HERB's
    answerable questions (the ids with ::a::), each once."""
    sets, by_sha = {}, {}
    for name, file in SETS.items():
        path = ROOT / file
        by_sha[hashlib.sha256(path.read_bytes()).hexdigest()] = name
        sets[name] = [json.loads(x)["id"] for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
    every = [i for ids in sets.values() for i in ids]
    if len(every) != len(set(every)) or set(every) != {i for i in refs if "::a::" in i}:
        sys.exit(f"the question sets {list(SETS)} are not HERB's answerable questions, each once")
    return sets, by_sha


def runs_tables(arm, folder, refs, sets, by_sha):
    d = ROOT / "output/k=chars" / folder
    man, ev = read(d / "run_manifest.json"), read(d / "eval_manifest.json")
    env = environment(d, man)
    set_name = by_sha.get(man["inputs"]["ids_sha256"])
    if set_name is None:
        sys.exit(f"{folder}: the run's question set is none of {list(SETS)}")
    ids = sets[set_name]
    if (man["n_questions"], man["n_ran"], man["n_failed"]) != (len(ids), len(ids), 0):
        sys.exit(f"{folder}: not every one of the {len(ids)} questions of {set_name} was answered")
    outputs = rows(d / "arm_outputs.jsonl")
    if sorted(r["id"] for r in outputs) != sorted(ids):
        sys.exit(f"{folder}: the rows are not the questions of {set_name}, each once")
    cells = {}
    for r in rows(d / "eval_results.jsonl"):
        cells.setdefault(r["question_id"], {})[r["metric"]] = r
    calls = rows(d / "eval_calls.jsonl")
    judge, _ = judge_cost(folder, ev, calls, cells, d)
    judging_seconds(calls, set(cells), folder)
    # the run without answers over the index with the employee and customer records: which questions are now
    # handed other texts, and every question's search
    p = ROOT / "output/k=chars" / RETRIEVED[folder]
    pman = read(p / "run_manifest.json")
    if (pman["inputs"]["ids_sha256"] != man["inputs"]["ids_sha256"]
            or (pman["n_questions"], pman["n_ran"], pman["n_failed"]) != (len(ids), len(ids), 0)):
        sys.exit(f"{RETRIEVED[folder]}: not a run of every question of {set_name}")
    passed = {o["id"]: o for o in rows(p / "arm_outputs.jsonl")}
    if sorted(passed) != sorted(ids) or any(o.get("search_time_s") is None for o in passed.values()):
        sys.exit(f"{RETRIEVED[folder]}: not every question of {set_name} once, each with its search time")
    changed = {o["id"] for o in outputs if delivered(o) != delivered(passed[o["id"]])}
    keep = set(ids) - changed
    red = redone(arm)
    if {i for i in red["rows"] if i in set(ids)} != changed:
        sys.exit(f"{REDONE[arm]}: the questions of {set_name} it answered are not those now handed other texts")
    if any(delivered(red["rows"][i]) != delivered(passed[i]) for i in changed):
        sys.exit(f"{REDONE[arm]}: a question was handed other texts than in {RETRIEVED[folder]}")
    red["used"] |= changed

    def texts(os):
        """The texts RAGAS embeds for these questions' scores: each answer and each reference, as its samples
        hold them (eval/ragas.py _to_sample)."""
        return {t for o in os for t in (" ".join(str(g) for g in refs[o["id"]]["ground_truth"]), o["answer"] or "")}

    # the judging of the questions that keep the run's answer, from the run's records; the calls in judge_usage
    # that no record ties to a question stay with them. The others' judging is REDONE's
    seconds, failed_judging, judge_in, judge_out, embedded, counted = judging_of(
        calls, keep, texts(o for o in outputs if o["id"] in keep), folder)
    of_changed = texts(o for o in outputs if o["id"] in changed)
    if any(c["texts"][0] not in of_changed for n, c in enumerate(calls) if c.get("kind") == "embed" and n not in counted):
        sys.exit(f"{folder}: an embedding request of the scoring is of no question's text")
    recorded = judge_records.usage_of(calls)
    if judge["tokens_in"] < recorded["tokens_in"] or judge["tokens_out"] < recorded["tokens_out"]:
        sys.exit(f"{folder}: judge_usage holds fewer tokens than the call records")
    judge_in += judge["tokens_in"] - recorded["tokens_in"]
    judge_out += judge["tokens_out"] - recorded["tokens_out"]
    r_seconds, r_failed, r_in, r_out, r_embedded, r_counted = judging_of(
        red["calls"], changed, texts(red["rows"][i] for i in changed), REDONE[arm], red["claimed"])
    red["claimed"] |= r_counted
    added, added_seconds, added_in, added_out, added_failed = rejudged(arm, d, cells, ev["judge_model"], set_name,
                                                                       keep)
    for qid, scores in added.items():
        cells[qid].update(scores)
    # the rows: the run's where the texts handed over are the same, REDONE's where they changed
    final_rows = {o["id"]: o for o in outputs}
    outputs = [red["rows"][o["id"]] if o["id"] in changed else o for o in outputs]
    for i in changed:
        cells[i] = red["cells"][i]
    source = {o["id"]: o for o in rows(ROOT / "output/k=chars" / MADE_IN[folder] / "arm_outputs.jsonl")} \
        if folder in MADE_IN else {}
    embedding_source = {o["id"]: o for o in rows(ROOT / "output/k=chars" / EMBEDDED_IN[folder] / "arm_outputs.jsonl")} \
        if folder in EMBEDDED_IN else {}
    # the source's vectors against the run's own, text by text: every text compared, none below the cosine
    vectors = read(ROOT / "output/k=chars" / EMBEDDED_IN[folder] / VECTOR_CHECK) if folder in EMBEDDED_IN else {}
    # the run's rows of the questions answered again as the run recorded them, before the corrections made after
    # it: the store keeps a first call's record as it was made
    p = d / "arm_outputs.as_recorded.jsonl"
    as_recorded = {o["id"]: o for o in rows(p) if o["id"] in changed} if changed and p.exists() else {}
    questions = []
    for r in outputs:
        again = r["id"] in changed
        own = {}
        for e in r["retrieval"]:
            if e["kind"] in own:
                sys.exit(f"{arm} {r['id']}: two {e['kind']} models in the arm's own calls")
            own[e["kind"]] = e
        arm_calls = [c for c in r["calls"] if c.get("role") != "answer"]
        # entries a REDONE row recorded itself: the harness of that run writes a call's time as its answering
        # try's time alone (time_s = request_s), the gate's wait and the failed tries beside it, not in it
        fresh = {id(r["generator"])} | {id(e) for e in own.values()} if again else set()
        if again and (r.get("meta") or {}).get("interpreter") is not None:
            # REDONE read the interpreter's answer from the store of the final runs' first calls and asked no
            # model: the first call counts as it counts for the final run's row
            f = final_rows[r["id"]]
            # the row keeps a copy of the stored call's record, as corrected or as recorded; any other record is a
            # call made anew
            kept = [source[r["id"]]] if source else [f] + ([as_recorded[r["id"]]] if r["id"] in as_recorded else [])
            stored = [c for o in kept for c in o["calls"] if c.get("role") != "answer" and c.get("kind") == "chat"]
            if (any(c not in stored for c in arm_calls if c.get("kind") == "chat")
                    or interpreter_and_embeddings(r, vectors=False)[0] != interpreter_and_embeddings(f, vectors=False)[0]):
                sys.exit(f"{arm} {r['id']}: {REDONE[arm]} asked the interpreter again, or got another answer")
            arm_calls = [c for c in arm_calls if c.get("kind") != "chat"]
            own.pop("chat", None)
            if not source:
                own["chat"] = next(e for e in f["retrieval"] if e["kind"] == "chat")
                arm_calls += [c for c in f["calls"] if c.get("role") != "answer" and c.get("kind") == "chat"]
        records = [(r["generator"], [c for c in r["calls"] if c.get("role") == "answer"])] + [
            (e, [c for c in arm_calls if c.get("kind") == k]) for k, e in own.items()]
        if source:
            o = source[r["id"]]
            records += [(e, [c for c in o["calls"] if c.get("role") != "answer" and c.get("kind") == e["kind"]])
                        for e in o["retrieval"]]
            # a question answered again embedded its texts on another machine: its vectors are compared by text
            if "chat" in own or (interpreter_and_embeddings(r, vectors=not again)
                                 != interpreter_and_embeddings(o, vectors=not again)):
                sys.exit(f"{arm} {r['id']}: the interpreter's answers or the query embeddings are not those of "
                         f"{MADE_IN[folder]}")
            made = {e["kind"]: e for e in o["retrieval"]}
            # this run's own embedding calls, for texts its machine had not embedded before, give way to
            # the calls that embedded all of them
            own["chat"], own["embed"] = made["chat"], made["embed"]
        if embedding_source:
            o = embedding_source[r["id"]]
            if interpreter_and_embeddings(r, vectors=False) != interpreter_and_embeddings(o, vectors=False):
                sys.exit(f"{arm} {r['id']}: the interpreter's answers or the texts embedded are not those of "
                         f"{EMBEDDED_IN[folder]}")
            v = vectors["per_question"].get(r["id"])
            if not v or v["compared"] != v["texts"] or v["min_cosine"] < vectors["min_cosine_required"]:
                sys.exit(f"{arm} {r['id']}: not every vector of {EMBEDDED_IN[folder]} is the run's own by its "
                         f"cosine ({VECTOR_CHECK})")
            made = {e["kind"]: e for e in o["retrieval"]}
            if "embed" not in made:
                sys.exit(f"{arm} {r['id']}: {EMBEDDED_IN[folder]} embedded nothing for the question")
            records += [(made["embed"], [c for c in o["calls"] if c.get("role") != "answer" and c.get("kind") == "embed"])]
            # the run's own embedding calls, for texts its machine had not embedded before, give way to the
            # calls that embedded all of them; the interpreter stays the run's own
            own["embed"] = made["embed"]
        mine = cells[r["id"]]
        if set(mine) != {s for s, _, _ in SCORES} or any(c["status"] != "ok" for c in mine.values()):
            sys.exit(f"{arm} {r['id']}: not every score is there and ok")
        g, interp, qe = r["generator"], own.get("chat"), own.get("embed")
        # a model call's time is its answering try; the seconds of failed tries are taken out of
        # the call, the question and the run (his 2026-10-21 "correct this"). What is taken out,
        # retry_s, must be the seconds of the tries that did not answer and of the pauses after tries,
        # as the call records hold them (two clock readings of the same tries: within 1e-4 s), and
        # the call's time must hold its tries and its wait; the rest of it is the harness's own time
        # in the call and stays in it, as in the final runs (the main session's decision of 2026-10-25)
        for call, recs in records:
            if (abs(call["retry_s"] - failed_tries_seconds(recs)) > 1e-4
                    or (abs(call["time_s"] - call["request_s"]) > 1e-6 if id(call) in fresh
                        else call["time_s"] - call["request_s"] - call["wait_s"] - call["retry_s"] < -1e-6)):
                sys.exit(f"{arm} {r['id']}: a call's retry_s is not its records' failed tries and pauses, or "
                         f"its time does not hold its tries and its wait")
        # the question's time is its steps' times added (his 2026-10-22 "THE RUN TIME IS THE ADDED TIME FOR
        # ALL THE THINGS IN THE FUCKING RUN"), each step as its column holds it: an interpreter call or
        # embedding calls made in MADE_IN's or EMBEDDED_IN's run, or read from the run's store (its record
        # given back by correction store-hit-calls/1), by their answering try, in place of the run's own;
        # the harness's own time between the steps is in no step
        # the search over the index that holds the employee and customer records, from RETRIEVED: one method,
        # one machine and one index for every question and arm
        parts = {"search_seconds": passed[r["id"]]["search_time_s"],
                 "interpreter_seconds": interp and (interp["request_s"] if id(interp) in fresh
                                                    else interp["time_s"] - interp["retry_s"]),
                 "query_embedding_seconds": qe and (qe["request_s"] if id(qe) in fresh
                                                    else qe["time_s"] - qe["retry_s"] - qe["wait_s"]),
                 "answer_seconds": g["request_s"] if id(g) in fresh else g["time_s"] - g["retry_s"]}
        # the characters the cut handed over, checked against the texts themselves
        chars = r["meta"]["char_budget"]["chars"]
        if chars != sum(len(t) for t in r["contexts"]) or chars > man["char_budget"]:
            sys.exit(f"{arm} {r['id']}: the characters handed over are not those of the retrieved context")
        ref = refs[r["id"]]
        if ref["question"] != r["question"]:
            sys.exit(f"{arm} {r['id']}: the question is not the questions file's")
        qtype = one((c["type"] for c in mine.values()), f"{r['id']} type")
        per_id = dict.fromkeys(("string_presence_per_id", "string_presence_per_id_precision",
                                "string_presence_per_id_f1"))
        if qtype in ID_TYPES:
            items = ref["ground_truth"]
            if not all(ids_and_links(str(item)) for item in items):
                sys.exit(f"{arm} {r['id']}: an item of HERB's reference list holds no ID or link of the forms")
            recall = string_presence_per_id(r["answer"], items)
            precision = string_presence_per_id_precision(r["answer"], items)
            # an answer that gives no ID or link scores 0, as in HERB's set F1
            per_id = {"string_presence_per_id": recall, "string_presence_per_id_precision": precision,
                      "string_presence_per_id_f1": (2 * precision * recall / (precision + recall)
                                                    if precision is not None and precision + recall else 0.0)}
        questions.append({
            "arm": arm,
            "question_id": r["id"],
            "question_type": qtype,
            "question": r["question"],
            "reference_answer": " ".join(str(g) for g in ref["ground_truth"]),
            "reference_context_ids": ", ".join(str(c) for c in ref["citations"]),
            "answer": r["answer"],
            **{s: mine[s]["value"] for s, _, _ in SCORES},
            **per_id,
            **parts,
            "total_seconds": sum(v for v in parts.values() if v),
            "interpreter_tokens_in": interp and interp["tokens_in"],
            "interpreter_tokens_out": interp and interp["tokens_out"],
            "query_embedding_tokens": qe and qe["tokens_in"],
            "retrieved_context_characters": chars,
            "answer_tokens_in": g["tokens_in"],
            "answer_tokens_out": g["tokens_out"],
        })
    setup = {
        "dataset": DATASET, "dataset_licence": LICENCE,
        "character_limit": man["char_budget"], "answer_model": model_name(man["generator_model"]),
        "judge_model": model_name(ev["judge_model"]), "embedding_model": model_name(env["embedder"]["model"]),
    }
    # his 2026-10-21 "cba adding interpreter, it's just the answer model"
    if man["interpreter_model"] is not None and model_name(man["interpreter_model"]) != setup["answer_model"]:
        sys.exit(f"{arm}: the interpreter's model is not the answer model")
    # a run's time, and its judging's, is the time of all the things in it added: its questions' times,
    # never the clock of the run as a whole (his 2026-10-22 "THE RUN TIME IS THE ADDED TIME FOR ALL THE
    # THINGS IN THE FUCKING RUN", 2026-10-23 "wall time? this is not a thing")
    # one setup for every row: the redone and the retrieved run's as the final run's
    rman = red["man"]
    if ((rman["char_budget"], model_name(rman["generator_model"]), model_name(red["ev"]["judge_model"]),
         model_name(rman["environment"]["embedder"]["model"]))
            != (setup["character_limit"], setup["answer_model"], setup["judge_model"], setup["embedding_model"])
            or (pman["char_budget"], model_name(pman["environment"]["embedder"]["model"]))
            != (setup["character_limit"], setup["embedding_model"])):
        sys.exit(f"{arm}: {REDONE[arm]} or {RETRIEVED[folder]} ran with another setup than {folder}")
    run = {
        "run_seconds": sum(q["total_seconds"] for q in questions),
        "judge_tokens_in": judge_in + r_in + added_in, "judge_tokens_out": judge_out + r_out + added_out,
        "judge_embedding_tokens": embedded + r_embedded,
        "judging_seconds": seconds + r_seconds + added_seconds,
    }
    failed_answer = sum(1 for r in outputs for c in r["calls"] if c.get("role") == "answer"
                        for t in c.get("attempts") or [] if t.get("outcome") != "ok")
    print(f"{folder}: {set_name}, {len(questions)} rows, {len(changed)} of them answered and judged again in "
          f"{REDONE[arm]}; answer calls' tries that did not answer {failed_answer}, "
          f"their seconds and pauses out: {sum(r['generator']['retry_s'] for r in outputs):.1f} s; seconds of judge "
          f"calls' tries that did not answer, and their pauses, left out of the judging: "
          f"{failed_judging + r_failed + added_failed:.1f} s")
    return questions, setup, run, set_name


def construction_table():
    """One row per property, one column per arm: what built the index that the final run read, one model
    at a time (his 2026-10-21 "The ONLY thing that matters, at all, is the iteration that ran the
    RAGAS"). For the artefact that is read from test/graph/build_cost.json: the descriptions and tags
    of the chunks in the graph, the vectors, the Opus judgements the facet network learned from, the
    encoding of every edge, and the steps with no model; work replaced later, trials and tests stay
    out. Every run of an arm must have read the same index, by what the index holds; the column holds
    the build as the arm's run on the 715 timed it."""
    # the index each arm reads now: that of its runs without answers over the corpus with the employee and
    # customer records (RETRIEVED), which its redone run (REDONE) must hold too
    now = {a: [RETRIEVED[f] for f in fs] for a, fs in RUNS.items()}
    of_runs = {a: [read(ROOT / "output/k=chars" / f / "run_manifest.json") for f in fs] for a, fs in now.items()}

    # an index is the same when it holds the same: the same units, their text and its digest, from the
    # same corpus, made by the same method or embedding model; not the same cache file or build seconds
    # (the main session's decision of 2026-10-25: the vector runs' two cache files hold the same units
    # and texts, their vectors equal within float rounding)
    def holds(arm, folder, m):
        i, emb = m["index"], environment(ROOT / "output/k=chars" / folder, m)["embedder"]
        out = {"corpus.sha256": m["inputs"]["corpus"]["sha256"]}
        embedder = {f"embedder.{k}": emb.get(k) for k in ("model", "revision", "dtype")}
        if arm == "lucene":
            out.update({f"index.{k}": i.get(k) for k in ("units", "text_chars", "unit_text", "units_sha256",
                                                          "method", "k1", "b", "stopwords", "stemmer", "bm25s",
                                                          "PyStemmer")})
        elif arm == "vector":
            out.update({f"index.{k}": i.get(k) for k in ("units", "text_chars", "unit_text", "units_sha256",
                                                          "matrix_shape", "matrix_dtype", "matrix_bytes")},
                       **embedder)
        else:
            out.update({"index.graph": i.get("graph"), "index.facet_layer_run": i.get("facet_layer_run"),
                        "index.counts": json.dumps(i.get("counts"), sort_keys=True),
                        "index.build_cost_file_sha256": i.get("build_cost_file_sha256"),
                        "graph": json.dumps(m.get("graph"), sort_keys=True)}, **embedder)
        return out

    unequal = []
    for a, ms in of_runs.items():
        names = now[a] + [REDONE[a]]
        held = [holds(a, f, m) for f, m in zip(names, ms + [read(ROOT / "output/k=chars" / REDONE[a] / "run_manifest.json")])]
        unequal += [f"{a} {k}: " + ", ".join(f"{f} {h[k]}" for f, h in zip(names, held))
                    for k in held[0] if any(h[k] != held[0][k] for h in held)]
    if unequal:
        sys.exit("the runs of an arm read different indexes, by what the indexes hold:\n  " + "\n  ".join(unequal))
    on_715 = next(hashlib.sha256((ROOT / f).read_bytes()).hexdigest() for s, f in SETS.items() if s == "answerable715")
    mans = {}
    for a, ms in of_runs.items():
        # the build of the index that the arm's run on the 715 read: one timer covered it (the vector index the
        # run on the hundred read holds the same units, but the time of its first part was read off its parts'
        # write times); the artefact's graph is the same for both runs
        on = [m for m in ms if m["inputs"]["ids_sha256"] == on_715]
        if len(on) != 1:
            sys.exit(f"{a}: not one run on answerable715, whose index build the construction holds")
        mans[a] = on[0]
        # for the record: each run's own build, of which only the 715's goes into the column
        print(f"{a}: build seconds of each run's index: " + ", ".join(
            f"{f} {m['build_stats'].get('build_time_s')}" for f, m in zip(now[a], ms)))
    cost = read(ROOT / "test/graph/build_cost.json")
    steps = {s["id"]: s for s in cost["steps"]}

    def left_out(step, starts):
        hits = [x["values"] for x in cost["left_out"] if x["step"] == step and x["what"].startswith(starts)]
        if len(hits) != 1:
            sys.exit(f"build_cost.json left_out: not one entry {step!r} {starts!r}")
        return hits[0]

    tagging = steps["descriptions_and_tags"]
    graph_chunks = tagging["the_4869_chunks_of_this_graph"]["extract"]
    # a model's time is its calls' times added, as a run's (his 2026-10-22 "THE RUN TIME IS THE ADDED TIME
    # FOR ALL THE THINGS IN THE FUCKING RUN"), never the clock of a pass that ran calls at once, each call by
    # its answering try (his 2026-10-21 "correct this"), recounted from the tagging run's own records
    haiku = read(ROOT / "test/graph/build_cost_haiku_answering_tries.json")
    left = haiku["answers_the_graph_does_not_hold"]
    if (haiku["as_build_cost_json_had_it"] != graph_chunks["seconds_summed_per_call"]
            or (haiku["calls"] + left["calls"], haiku["tokens_in"] + left["tokens_in"], haiku["tokens_out"] + left["tokens_out"])
            != (graph_chunks["calls_with_usage"], graph_chunks["tokens_in"], graph_chunks["tokens_out"])):
        sys.exit("build_cost_haiku_answering_tries.json is not the recount of build_cost.json's haiku calls")
    haiku_seconds = haiku["seconds"]
    vectors = steps["embeddings"]["parts"]
    runs = left_out("facet_layer", "the Opus judge runs one by one")
    learned_from = [runs["control"], runs["train"]]    # calls, in, out, summed, elapsed; the held-out only tested
    facet = steps["facet_layer"]["parts"]
    no_model = (sum(steps["graph_copies"]["seconds"].values())
                + vectors["write_tag_vectors"]["seconds_elapsed"] + vectors["write_description_vectors"]["seconds_elapsed"]
                # round 1 without its tests: the probe that chose and fitted the head, and the scoring; its held-out
                # bootstrap and its diagnostics built nothing the final runs read
                + left_out("facet_layer", "parts inside round 1")["probe_s"]
                + left_out("facet_layer", "parts inside round 1")["scoring_s"]
                + facet["bootstrap_refits"]["seconds_elapsed"]
                + facet["write"]["seconds_elapsed"]
                + sum(steps["later_edges_and_fields"]["seconds"].values())
                # the employee and customer chunks' facet layer and edges (2026-10-31), but their encoding
                + sum(v for k, v in steps["metadata_facet_layer"]["seconds"].items() if k != "encode_574_edges")
                + sum(steps["metadata_person_and_customer_edges"]["seconds"].values()))
    # the 63 readable tag names were embedded on the laptop's CPU; their time is that of the same requests
    # on the desktop's graphics card, where the other two parts ran (his 2026-10-27 "embeddings via gpu")
    readable = read(ROOT / "test/graph/build_cost_readable_names_gpu_times.json")
    part = vectors["tags_with_underscores_in_readable_form"]
    if (readable["measured"]["device"] != "cuda" or len(readable["requests"]) != part["texts"]
            or sum(x["tokens_in"] for x in readable["requests"]) != part["tokens_in"]):
        sys.exit("build_cost_readable_names_gpu_times.json is not the readable names' requests on the graphics card")
    # the raw names' vectors were replaced in the graph by the readable names': their share of the tag pass, by the
    # pass's own rate, is not part of what built the index (his 2026-10-21 "The ONLY thing that matters, at all, is
    # the iteration that ran the RAGAS")
    names = read(ROOT / "test/graph/build_cost_readable_names_tokens.json")
    if ((names["tags"], names["tags_tokens"], names["raw_names"], names["readable_names_tokens"])
            != (vectors["tags"]["texts"], vectors["tags"]["tokens_in"], part["texts"], part["tokens_in"])):
        sys.exit("build_cost_readable_names_tokens.json is not the count of build_cost.json's tag pass")
    kept_tags = vectors["tags"]["texts"] - names["raw_names"]
    embed_texts = kept_tags + part["texts"] + vectors["descriptions"]["texts"]
    embed_tokens = (vectors["tags"]["tokens_in"] - names["raw_names_tokens"] + part["tokens_in"]
                    + vectors["descriptions"]["tokens_in"])
    embed_seconds = (vectors["tags"]["seconds_elapsed"] * kept_tags / vectors["tags"]["texts"]
                     + vectors["descriptions"]["seconds_elapsed"] + sum(x["encode_s"] for x in readable["requests"]))
    # the Opus calls' records keep only the clock around each call; their CLI transcripts keep the prompt and the
    # reply, which on the final runs' calls give the CLI's duration_api_ms within a median of 0.02 s
    opus = read(ROOT / "test/graph/build_cost_opus_transcript_times.json")
    if [opus["sets"][s]["calls"] for s in ("control", "train")] != [r[0] for r in learned_from]:
        sys.exit("build_cost_opus_transcript_times.json is not the Opus calls of the control and train runs")
    opus_seconds = opus["seconds"]
    encoder_seconds = (facet["encoder_cache"]["seconds_elapsed"]
                       + steps["metadata_facet_layer"]["seconds"]["encode_574_edges"])
    vector = mans["vector"]["build_stats"]["models"][0]
    embedder = model_name(vector["model"])
    if any(model_name(environment(ROOT / "output/k=chars" / f, m)["embedder"]["model"]) != embedder
           for f, m in zip(now["artefact"], of_runs["artefact"])):
        sys.exit("the arms name different embedding models")
    encoder = model_name(facet["encoder_cache"]["model"].split(", ")[0])
    a = "artefact"
    things = [
        ("index", "Index type.", "", {k: INDEX_KIND[k][0] for k in RUNS}),
        ("unit", "Unit stored in the index.", "", {k: INDEX_KIND[k][1] for k in RUNS}),
        ("units in the index", "Units in the index. Artefact: the chunks in the graph; the arm retrieves from all of "
                               "them but the chunks of one of HERB's metadata files, salesforce_team.json.", "count",
         {"lucene": mans["lucene"]["index"]["units"], "vector": mans["vector"]["index"]["units"],
          a: mans[a]["graph"]["census"]["labels"]["Chunk"]}),
        ("build time, seconds", "Time to build the index that the arm's runs read. Artefact: the times of "
                                "the build steps added: the claude-haiku-4-5 and claude-opus-5 steps as the times of "
                                "their calls added, the embedding as in its row, the other steps as their recorded "
                                "times; the chunking and loading of the chunks has no recorded time and is not in "
                                "it.", "seconds",
         {"lucene": mans["lucene"]["build_stats"]["build_time_s"], "vector": mans["vector"]["build_stats"]["build_time_s"],
          a: haiku_seconds + embed_seconds + opus_seconds + encoder_seconds + no_model}),
        (f"{embedder}, texts embedded", "Texts embedded for the index, as the index holds them.", "count",
         {"vector": vector["calls"], a: embed_texts}),
        (f"{embedder}, tokens", "Tokens embedded for the index, counted on the text with its mode's prefix and the "
                                "tokenizer's special tokens.", "tokens",
         {"vector": vector["tokens_in"], a: embed_tokens}),
        (f"{embedder}, seconds", "Embedding time for the index, on the graphics card. Vector: its requests' times "
                                 "added. Artefact: the elapsed times of its two passes (model load and saving "
                                 "included), less the share of the 63 tags whose vectors were made again from their "
                                 "readable names, by the pass's rate, and those 63 requests' times.", "seconds",
         {"vector": vector["seconds_elapsed"], a: embed_seconds}),
        ("claude-haiku-4-5, calls", "Calls to claude-haiku-4-5 whose descriptions and tags the index holds.", "count",
         {a: haiku["calls"]}),
        ("claude-haiku-4-5, tokens in", "Input tokens of those calls.", "tokens", {a: haiku["tokens_in"]}),
        ("claude-haiku-4-5, tokens out", "Output tokens of those calls.", "tokens", {a: haiku["tokens_out"]}),
        ("claude-haiku-4-5, seconds", "Time of those calls added, each by its answering try as the tagging run "
                                      "recorded it; a call answered on a later try, whose record also holds its failed "
                                      "tries, by a + b x its output tokens fitted on the calls answered on the first "
                                      "try.", "seconds", {a: haiku_seconds}),
        ("claude-opus-5, calls", "Calls to claude-opus-5 for the index.", "count",
         {a: sum(r[0] for r in learned_from)}),
        ("claude-opus-5, tokens in", "Input tokens of those calls.", "tokens", {a: sum(r[1] for r in learned_from)}),
        ("claude-opus-5, tokens out", "Output tokens of those calls.", "tokens", {a: sum(r[2] for r in learned_from)}),
        ("claude-opus-5, seconds", "Time of those calls added, each from its prompt to its reply in the Claude CLI's "
                                   "own transcript of the call.", "seconds",
         {a: opus_seconds}),
        (f"{encoder}, seconds", f"Encoding time of {encoder} for the index.", "seconds", {a: encoder_seconds}),
    ]
    table = [{"property": t, **{k: values.get(k) for k in RUNS}} for t, _, _, values in things]
    return table, [{"sheet": CONSTRUCTION, "field": t, "meaning": m, "unit": u} for t, m, u, _ in things]


WIDE = {"question": 60, "reference_answer": 60, "answer": 80, "meaning": 100, "property": 50}


def workbook(path, title, sheets):
    """sheets: (name, [(heads, rows), ...]); the tables of one sheet stand under each other, a blank
    row between them."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title, wb.properties.creator = title, ""
    for name, blocks in sheets:
        ws = wb.create_sheet(name)
        widths = {}
        for n, (heads, table) in enumerate(blocks):
            if n:
                ws.append([])
            ws.append(heads)
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
            for row in table:
                if set(row) != set(heads):
                    sys.exit(f"{path.name} {name}: a row's columns differ from its table's")
                ws.append([row[h] for h in heads])
            for i, h in enumerate(heads, 1):
                longest = max([len(h)] + [len(str(r[h])) for r in table[:200] if r[h] is not None])
                widths[i] = max(widths.get(i, 10), min(WIDE.get(h, 40), longest + 2))
        if len(blocks) == 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for i, w in widths.items():
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    wb.save(path)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", default=str(ROOT / "finished_data"))
    out = Path(p.parse_args().out)
    out.mkdir(parents=True, exist_ok=True)
    questions, setups, runs, covered = [], {}, {arm: [] for arm in RUNS}, {arm: {} for arm in RUNS}
    refs = references()
    sets, by_sha = question_sets(refs)
    for arm, folders in RUNS.items():
        for folder in folders:
            rows_of_run, setups[folder], run_of, set_name = runs_tables(arm, folder, refs, sets, by_sha)
            if set_name in covered[arm]:
                sys.exit(f"{arm}: two runs on {set_name}")
            covered[arm][set_name] = len(rows_of_run)
            questions += rows_of_run
            runs[arm].append(run_of)
    # every question an arm's REDONE run answered is one of its sets' questions now handed other texts, and
    # each of that run's embedding requests is counted once
    for arm in RUNS:
        red = redone(arm)
        if (red["used"] != set(red["rows"])
                or red["claimed"] != {n for n, c in enumerate(red["calls"]) if c.get("kind") == "embed"}):
            sys.exit(f"{REDONE[arm]}: a question it answered, or an embedding request it made, is in no set's rows")
    # which question sets each arm's runs cover; an arm without a run on a set is named, and nothing
    # is decided about it here
    for arm in RUNS:
        lacks = [s for s in SETS if s not in covered[arm]]
        print(f"{arm}: " + ", ".join(f"{s} {n} questions" for s, n in covered[arm].items())
              + (f"; no run on {', '.join(lacks)}" if lacks else ""))
    if len({tuple(sorted(c)) for c in covered.values()}) > 1:
        print("the arms' rows, their run data and their means cover different questions")
    setup = next(iter(setups.values()))
    if any(s != setup for s in setups.values()):
        sys.exit("the runs differ in the setup; this export writes it once")
    # run data: what is the same for the three arms once, then a table with a column per arm, each
    # value the arm's runs' added
    same = [{"property": c, "value": setup[c]} for c, _, _ in SETUP]
    run = [{"property": c, **{arm: sum(r[c] for r in runs[arm]) for arm in RUNS}} for c, _, _ in RUN]
    # the mean of each measure of the results over an arm's questions in results; empty where the arm has none
    measures = [c for c, _, u in QUESTIONS if u]
    means = [{"mean of": c, **{arm: (lambda v: sum(v) / len(v) if v else None)(
        [q[c] for q in questions if q["arm"] == arm and q[c] is not None]) for arm in RUNS}} for c in measures]
    construction, construction_key = construction_table()
    key = ([{"sheet": PER_QUESTION, "field": c, "meaning": m, "unit": u} for c, m, u in QUESTIONS]
           + [{"sheet": PER_RUN, "field": c, "meaning": m, "unit": u} for c, m, u in SETUP + RUN if m]
           + [{"sheet": PER_RUN, "field": "mean of", "meaning": f"Mean over the arm's questions in {PER_QUESTION} "
                                                               f"of the field of {PER_QUESTION} named in the row; empty where "
                                                               "the arm has no value of the field.",
               "unit": "the unit of the field named in the row"}]
           + construction_key)
    sheets = [(PER_QUESTION, [([c for c, _, _ in QUESTIONS], questions)]),
              (PER_RUN, [(["property", "value"], same), (["property", *RUNS], run),
                         (["mean of", *RUNS], means)]),
              (CONSTRUCTION, [(["property", *RUNS], construction)]),
              (KEY, [(["sheet", "field", "meaning", "unit"], key)])]
    path = out / "full-grag-data.xlsx"
    workbook(path, "grag data", sheets)
    print(f"{path.name}: {path.stat().st_size:,} bytes, sheets {[n for n, _ in sheets]}, "
          f"rows {[[len(t) for _, t in blocks] for _, blocks in sheets]}")


if __name__ == "__main__":
    main()
